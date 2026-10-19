import os
import sys
import re
import json
import argparse
from datetime import datetime

# Configure output encoding for Windows terminal
sys.stdout.reconfigure(encoding='utf-8')

TRACKER_JSON_PATH = os.path.join('data', 'manual_test_tracker.json')
EXCEL_PATH = 'Curriculum_Manual_Test_Tracker.xlsx'
CSV_PATH = 'Curriculum_Manual_Test_Tracker.csv'
MD_PATH = 'CURRICULUM_TEST_TRACKER.md'
CHAPTERS_DATA_PATH = os.path.join('src', 'data', 'chaptersData.ts')

def extract_field(block, field_name):
    # Matches field: 'value' or field: "value" or field: value
    m = re.search(rf"{field_name}:\s*(?:'([^']*)'|\"([^\"]*)\"|([^\s,}}]+))", block)
    if m:
        return (m.group(1) if m.group(1) is not None else (m.group(2) if m.group(2) is not None else m.group(3)))
    return ""

def parse_chapters_data():
    """Extract all chapters from chaptersData.ts"""
    if not os.path.exists(CHAPTERS_DATA_PATH):
        raise FileNotFoundError(f"Cannot find {CHAPTERS_DATA_PATH}")

    with open(CHAPTERS_DATA_PATH, 'r', encoding='utf-8') as f:
        text = f.read()

    # Find each chapter object block { id: ..., number: ..., ... }
    blocks = re.findall(r"\{\s*id:\s*['\"][^'\"]+['\"].*?\}", text, re.DOTALL)
    chapters = []
    for b in blocks:
        cid = extract_field(b, "id")
        num = extract_field(b, "number")
        title = extract_field(b, "title")
        grade = extract_field(b, "grade")
        subject = extract_field(b, "subject")
        icon = extract_field(b, "icon")
        summary = extract_field(b, "conceptSummary")
        if not (cid and num and title and grade and subject):
            continue

        grade_str = str(grade).strip()
        if grade_str.isdigit():
            grade_display = f"Class {grade_str}"
            grade_sort_key = int(grade_str) + 2  # LKG=1, UKG=2, Class 1=3...
        elif grade_str.upper() == 'LKG':
            grade_display = "LKG (Balvatika 2)"
            grade_sort_key = 1
        elif grade_str.upper() == 'UKG':
            grade_display = "UKG (Balvatika 3)"
            grade_sort_key = 2
        else:
            grade_display = grade_str
            grade_sort_key = 99

        chapters.append({
            "id": cid,
            "chapterNumber": int(num),
            "title": title,
            "grade": grade_str,
            "gradeDisplay": grade_display,
            "gradeSortKey": grade_sort_key,
            "subject": subject.lower(),
            "subjectDisplay": "Maths" if subject.lower() == 'maths' else ("English" if subject.lower() == 'english' else "Hindi"),
            "icon": icon,
            "conceptSummary": summary,
        })

    # Sort chapters logically: Grade sort key, then Subject (Maths, English, Hindi), then chapterNumber
    subject_order = {"maths": 1, "english": 2, "hindi": 3}
    chapters.sort(key=lambda x: (x["gradeSortKey"], subject_order.get(x["subject"], 9), x["chapterNumber"]))
    return chapters

def load_or_init_tracker():
    """Load existing JSON database or initialize from chaptersData.ts"""
    os.makedirs('data', exist_ok=True)
    existing_map = {}

    if os.path.exists(TRACKER_JSON_PATH):
        try:
            with open(TRACKER_JSON_PATH, 'r', encoding='utf-8') as f:
                data = json.load(f)
                for item in data.get('chapters', []):
                    existing_map[item['id']] = item
        except Exception as e:
            print(f"Warning: could not read existing tracker JSON: {e}")

    chapters_catalog = parse_chapters_data()
    merged = []

    for ch in chapters_catalog:
        cid = ch['id']
        if cid in existing_map:
            ex = existing_map[cid]
            # Preserve existing test history while updating catalog metadata
            ch['status'] = ex.get('status', 'Pending Manual Test')
            ch['firstTestedDate'] = ex.get('firstTestedDate', None)
            ch['revisionDates'] = ex.get('revisionDates', [])
            ch['totalRevisions'] = ex.get('totalRevisions', len(ex.get('revisionDates', [])))
            ch['issuesFixed'] = ex.get('issuesFixed', '')
            ch['lastAppVersion'] = ex.get('lastAppVersion', '')
            ch['notes'] = ex.get('notes', '')
        else:
            ch['status'] = 'Pending Manual Test'
            ch['firstTestedDate'] = None
            ch['revisionDates'] = []
            ch['totalRevisions'] = 0
            ch['issuesFixed'] = ''
            ch['lastAppVersion'] = ''
            ch['notes'] = ''
        merged.append(ch)

    return merged

def save_tracker_json(chapters):
    payload = {
        "updatedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "totalChapters": len(chapters),
        "totalTested": sum(1 for c in chapters if c.get('firstTestedDate')),
        "chapters": chapters
    }
    with open(TRACKER_JSON_PATH, 'w', encoding='utf-8') as f:
        json.dump(payload, f, ensure_ascii=False, indent=2)
    return payload

def record_lesson_update(grade, subject, chapter_num, issues, test_date=None, version="v1.3.18 (Build 37)", notes=""):
    """Update a specific lesson with new test/revision details"""
    chapters = load_or_init_tracker()
    if not test_date:
        test_date = datetime.now().strftime("%Y-%m-%d")

    grade_normalized = str(grade).strip()
    if grade_normalized.lower().startswith('class'):
        grade_normalized = grade_normalized.lower().replace('class', '').strip()
    grade_normalized = grade_normalized.upper() if grade_normalized.upper() in ['LKG', 'UKG'] else grade_normalized

    subject_normalized = subject.lower().strip()
    chapter_num = int(chapter_num)

    found = False
    for ch in chapters:
        if ch['grade'].upper() == grade_normalized and ch['subject'] == subject_normalized and ch['chapterNumber'] == chapter_num:
            found = True
            is_first_time = ch['firstTestedDate'] is None

            if is_first_time:
                ch['firstTestedDate'] = test_date
                ch['revisionDates'] = [test_date]
                ch['totalRevisions'] = 1
                ch['issuesFixed'] = f"[{test_date}] {issues}"
            else:
                # Subsequent update / Revision
                ch['revisionDates'].append(test_date)
                ch['totalRevisions'] += 1
                if ch['issuesFixed']:
                    ch['issuesFixed'] += f" | [{test_date} Rev #{ch['totalRevisions']}] {issues}"
                else:
                    ch['issuesFixed'] = f"[{test_date} Rev #{ch['totalRevisions']}] {issues}"

            ch['status'] = 'Tested & Updated'
            ch['lastAppVersion'] = version
            if notes:
                ch['notes'] = (ch.get('notes', '') + " " + notes).strip()

            print(f"✅ Successfully updated {ch['gradeDisplay']} | {ch['subjectDisplay']} | Ch {ch['chapterNumber']}: '{ch['title']}'")
            print(f"   First Tested Date: {ch['firstTestedDate']}")
            print(f"   All Dates: {', '.join(ch['revisionDates'])} (Total Revisions: {ch['totalRevisions']})")
            print(f"   Status: {ch['status']}")
            break

    if not found:
        print(f"❌ Error: Lesson not found for Grade '{grade}', Subject '{subject}', Chapter '{chapter_num}'")
        return False

    save_tracker_json(chapters)
    export_all(chapters)
    return True

def export_csv(chapters):
    import csv
    with open(CSV_PATH, 'w', encoding='utf-8-sig', newline='') as f:
        writer = csv.writer(f)
        writer.writerow([
            "S.No",
            "Class / Grade",
            "Subject",
            "Chapter No.",
            "Chapter Title",
            "Status",
            "First Tested Date",
            "Revision Dates",
            "Total Updates",
            "Issues Found & Fixed",
            "App Version",
            "Concept Summary"
        ])
        for idx, ch in enumerate(chapters, start=1):
            rev_dates_str = ", ".join(ch.get('revisionDates', []))
            writer.writerow([
                idx,
                ch['gradeDisplay'],
                ch['subjectDisplay'],
                f"Chapter {ch['chapterNumber']}",
                ch['title'],
                ch['status'],
                ch.get('firstTestedDate') or "-",
                rev_dates_str or "-",
                ch.get('totalRevisions', 0),
                ch.get('issuesFixed') or "No issues reported yet",
                ch.get('lastAppVersion') or "-",
                ch.get('conceptSummary') or ""
            ])
    print(f"📄 CSV generated: {CSV_PATH}")

def export_excel(chapters):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    # Sheet 1: Master Tracker
    ws_all = wb.active
    ws_all.title = "Master Tracker (All 167)"

    # Sheet 2: Tested & Updated Log
    ws_tested = wb.create_sheet(title="Tested & Updated Only")

    # Sheet 3: Progress Summary
    ws_summary = wb.create_sheet(title="Class Progress Summary")

    header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_fill_tested = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
    header_fill_summary = PatternFill(start_color="4C1D95", end_color="4C1D95", fill_type="solid")

    status_tested_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    status_tested_font = Font(name="Segoe UI", size=10, bold=True, color="166534")

    status_pending_fill = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid")
    status_pending_font = Font(name="Segoe UI", size=10, color="6B7280")

    regular_font = Font(name="Segoe UI", size=10)
    bold_font = Font(name="Segoe UI", size=10, bold=True)

    thin_border_side = Side(border_style="thin", color="E5E7EB")
    cell_border = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)

    headers = [
        "S.No",
        "Class / Grade",
        "Subject",
        "Chapter No.",
        "Chapter Title",
        "Status",
        "First Tested Date",
        "Revision Dates",
        "Total Updates",
        "Issues Found & Fixed",
        "App Version",
        "Concept Summary"
    ]

    # --- Populate Sheet 1 (All Chapters) ---
    ws_all.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws_all.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for idx, ch in enumerate(chapters, start=1):
        rev_dates = ", ".join(ch.get('revisionDates', [])) if ch.get('revisionDates') else "-"
        row_vals = [
            idx,
            ch['gradeDisplay'],
            ch['subjectDisplay'],
            f"Ch {ch['chapterNumber']}",
            f"{ch.get('icon', '')} {ch['title']}".strip(),
            ch['status'],
            ch.get('firstTestedDate') or "-",
            rev_dates,
            ch.get('totalRevisions', 0),
            ch.get('issuesFixed') or "Pending manual test",
            ch.get('lastAppVersion') or "-",
            ch.get('conceptSummary') or ""
        ]
        ws_all.append(row_vals)
        current_row = idx + 1

        for col_idx in range(1, len(headers) + 1):
            cell = ws_all.cell(row=current_row, column=col_idx)
            cell.font = regular_font
            cell.border = cell_border
            if col_idx in [1, 4, 7, 8, 9, 11]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [10, 12]:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            # Status column color
            if col_idx == 6:
                if ch['status'] == 'Tested & Updated':
                    cell.fill = status_tested_fill
                    cell.font = status_tested_font
                else:
                    cell.fill = status_pending_fill
                    cell.font = status_pending_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # Freeze header row on Sheet 1
    ws_all.freeze_panes = "A2"

    # --- Populate Sheet 2 (Tested & Updated Only) ---
    ws_tested.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws_tested.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill_tested
        cell.alignment = Alignment(horizontal="center", vertical="center")

    tested_chapters = [c for c in chapters if c.get('firstTestedDate')]
    tested_chapters.sort(key=lambda x: x.get('firstTestedDate', ''), reverse=True)

    for idx, ch in enumerate(tested_chapters, start=1):
        rev_dates = ", ".join(ch.get('revisionDates', []))
        row_vals = [
            idx,
            ch['gradeDisplay'],
            ch['subjectDisplay'],
            f"Ch {ch['chapterNumber']}",
            f"{ch.get('icon', '')} {ch['title']}".strip(),
            ch['status'],
            ch.get('firstTestedDate') or "-",
            rev_dates or "-",
            ch.get('totalRevisions', 0),
            ch.get('issuesFixed') or "-",
            ch.get('lastAppVersion') or "-",
            ch.get('conceptSummary') or ""
        ]
        ws_tested.append(row_vals)
        current_row = idx + 1
        for col_idx in range(1, len(headers) + 1):
            cell = ws_tested.cell(row=current_row, column=col_idx)
            cell.font = regular_font
            cell.border = cell_border
            cell.fill = status_tested_fill if col_idx == 6 else PatternFill(fill_type=None)
            if col_idx == 6:
                cell.font = status_tested_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [1, 4, 7, 8, 9, 11]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in [10, 12]:
                cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws_tested.freeze_panes = "A2"

    # --- Populate Sheet 3 (Class-wise Progress Summary) ---
    summary_headers = ["Class / Grade", "Total Chapters", "Tested & Updated", "Pending Test", "% Completed", "Progress Status"]
    ws_summary.append(summary_headers)
    for col_idx in range(1, len(summary_headers) + 1):
        cell = ws_summary.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill_summary
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Group by grade
    grade_order = ["LKG", "UKG", "1", "2", "3", "4", "5", "6"]
    grade_names = {
        "LKG": "LKG (Balvatika 2)",
        "UKG": "UKG (Balvatika 3)",
        "1": "Class 1",
        "2": "Class 2",
        "3": "Class 3",
        "4": "Class 4",
        "5": "Class 5",
        "6": "Class 6"
    }

    total_all = 0
    tested_all = 0

    for g in grade_order:
        g_chapters = [c for c in chapters if c['grade'].upper() == g]
        total_c = len(g_chapters)
        tested_c = sum(1 for c in g_chapters if c.get('firstTestedDate'))
        pending_c = total_c - tested_c
        pct = round((tested_c / total_c * 100), 1) if total_c > 0 else 0
        status_txt = "🚀 Completed" if pct == 100 else ("🟡 In Progress" if tested_c > 0 else "⚪ Not Started")

        total_all += total_c
        tested_all += tested_c

        ws_summary.append([
            grade_names.get(g, g),
            total_c,
            tested_c,
            pending_c,
            f"{pct}%",
            status_txt
        ])

    # Total Row
    total_pct = round((tested_all / total_all * 100), 1) if total_all > 0 else 0
    ws_summary.append([
        "OVERALL TOTAL",
        total_all,
        tested_all,
        total_all - tested_all,
        f"{total_pct}%",
        f"{tested_all}/{total_all} Tested"
    ])

    summary_last_row = len(grade_order) + 2
    for col_idx in range(1, len(summary_headers) + 1):
        cell = ws_summary.cell(row=summary_last_row, column=col_idx)
        cell.font = bold_font
        cell.border = cell_border
        cell.fill = PatternFill(start_color="E9D5FF", end_color="E9D5FF", fill_type="solid")

    for r in range(2, summary_last_row):
        for c in range(1, len(summary_headers) + 1):
            cell = ws_summary.cell(row=r, column=c)
            cell.font = regular_font
            cell.border = cell_border
            if c >= 2:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # Set column widths nicely on all sheets
    for ws in [ws_all, ws_tested]:
        column_widths = {
            1: 8,   # S.No
            2: 18,  # Grade
            3: 12,  # Subject
            4: 12,  # Ch No
            5: 35,  # Title
            6: 22,  # Status
            7: 18,  # First date
            8: 24,  # Rev dates
            9: 14,  # Updates count
            10: 45, # Issues
            11: 20, # Version
            12: 40  # Summary
        }
        for col_idx, width in column_widths.items():
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = width

    for col_idx in range(1, 7):
        ws_summary.column_dimensions[get_column_letter(col_idx)].width = 20

    wb.save(EXCEL_PATH)
    print(f"📊 Excel Workbook generated: {EXCEL_PATH}")

def export_markdown(chapters):
    total = len(chapters)
    tested = sum(1 for c in chapters if c.get('firstTestedDate'))
    pending = total - tested
    pct = round(tested / total * 100, 1)

    tested_list = [c for c in chapters if c.get('firstTestedDate')]
    tested_list.sort(key=lambda x: x.get('firstTestedDate', ''), reverse=True)

    lines = []
    lines.append("# 📋 ClassChamp Curriculum Manual Test & Revision Tracker")
    lines.append("")
    lines.append(f"> **Last Updated**: `{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}`  ")
    lines.append(f"> **Overall Progress**: **{tested}/{total} Chapters ({pct}%) Tested & Updated** | **{pending} Pending**  ")
    lines.append(f"> **Files**: [`{EXCEL_PATH}`]({EXCEL_PATH}) | [`{CSV_PATH}`]({CSV_PATH})")
    lines.append("")
    lines.append("---")
    lines.append("")
    lines.append("## 🏆 Recently Tested & Updated Lessons")
    lines.append("")
    if tested_list:
        lines.append("| S.No | Class / Grade | Subject | Chapter | Lesson Name | 1st Tested Date | Revision Dates | Revisions | Issues Fixed / Updates | Version |")
        lines.append("|:---:|:---|:---|:---:|:---|:---:|:---|:---:|:---|:---:|")
        for idx, ch in enumerate(tested_list, start=1):
            rev_str = ", ".join([f"`{d}`" for d in ch.get('revisionDates', [])])
            first_d = f"`{ch.get('firstTestedDate')}`"
            lines.append(f"| {idx} | **{ch['gradeDisplay']}** | {ch['subjectDisplay']} | Ch {ch['chapterNumber']} | {ch.get('icon', '')} {ch['title']} | {first_d} | {rev_str} | **{ch.get('totalRevisions', 1)}** | {ch.get('issuesFixed', '-')} | `{ch.get('lastAppVersion', '-')}` |")
    else:
        lines.append("*No lessons marked as tested yet.*")

    lines.append("")
    lines.append("---")
    lines.append("")
    lines.append("## 📊 Class-wise Test Completion Summary")
    lines.append("")
    lines.append("| Class / Grade | Total Chapters | Tested & Updated | Pending | Progress | Status |")
    lines.append("|:---|:---:|:---:|:---:|:---:|:---|")

    grade_order = ["LKG", "UKG", "1", "2", "3", "4", "5", "6"]
    grade_names = {
        "LKG": "LKG (Balvatika 2)",
        "UKG": "UKG (Balvatika 3)",
        "1": "Class 1",
        "2": "Class 2",
        "3": "Class 3",
        "4": "Class 4",
        "5": "Class 5",
        "6": "Class 6"
    }

    for g in grade_order:
        g_chapters = [c for c in chapters if c['grade'].upper() == g]
        total_c = len(g_chapters)
        tested_c = sum(1 for c in g_chapters if c.get('firstTestedDate'))
        pending_c = total_c - tested_c
        g_pct = round((tested_c / total_c * 100), 1) if total_c > 0 else 0
        bar = "🟩" * int(g_pct // 20) + "⬜" * (5 - int(g_pct // 20))
        lines.append(f"| **{grade_names.get(g, g)}** | {total_c} | {tested_c} | {pending_c} | {bar} {g_pct}% | {'🚀 Completed' if g_pct == 100 else ('🟡 In Progress' if tested_c > 0 else '⚪ Pending')} |")

    lines.append(f"| **TOTAL** | **{total}** | **{tested}** | **{pending}** | **{pct}%** | **{tested}/{total} Tested** |")
    lines.append("")

    with open(MD_PATH, 'w', encoding='utf-8') as f:
        f.write("\n".join(lines))
    print(f"📝 Markdown tracker generated: {MD_PATH}")

def export_all(chapters):
    export_csv(chapters)
    export_excel(chapters)
    export_markdown(chapters)

def main():
    parser = argparse.ArgumentParser(description="ClassChamp Curriculum Manual Test Tracker")
    subparsers = parser.add_subparsers(dest="command")

    # Command: generate
    subparsers.add_parser("generate", help="Generate Excel, CSV, and Markdown from existing tracker data")

    # Command: update
    update_parser = subparsers.add_parser("update", help="Record a manual test and fix for a lesson")
    update_parser.add_argument("--grade", required=True, help="LKG, UKG, 1, 2, 3, 4, 5, 6")
    update_parser.add_argument("--subject", required=True, help="maths, english, hindi")
    update_parser.add_argument("--chapter", required=True, type=int, help="Chapter number (1, 2, 3...)")
    update_parser.add_argument("--issues", required=True, help="Description of issues found and fixed")
    update_parser.add_argument("--date", default=None, help="Test date (YYYY-MM-DD, defaults to today)")
    update_parser.add_argument("--version", default="v1.3.18 (Build 37)", help="App version/build")
    update_parser.add_argument("--notes", default="", help="Optional notes")

    args = parser.parse_args()

    if args.command == "update":
        record_lesson_update(
            grade=args.grade,
            subject=args.subject,
            chapter_num=args.chapter,
            issues=args.issues,
            test_date=args.date,
            version=args.version,
            notes=args.notes
        )
    else:
        chapters = load_or_init_tracker()
        save_tracker_json(chapters)
        export_all(chapters)

if __name__ == '__main__':
    main()
